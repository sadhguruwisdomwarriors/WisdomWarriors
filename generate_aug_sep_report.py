import asyncio
import json
import ssl
import urllib.request
from collections import defaultdict
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.services.youtube_read_service import calculate_youtube_monthly_metrics

# SSL context
ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE

SUPABASE_URL = "https://juxvmzccqbroqcoleobp.supabase.co"
SERVICE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6Imp1eHZtemNjcWJyb3Fjb2xlb2JwIiwicm9sZSI6InNlcnZpY2Vfcm9sZSIsImlhdCI6MTc3NDY1MjE0OSwiZXhwIjoyMDkwMjI4MTQ5fQ.ayuLBYZbPhxv3DBUVvlcSX0y-0zlXSBl6jTBMkCWekE"
headers = {"apikey": SERVICE_KEY, "Authorization": f"Bearer {SERVICE_KEY}", "Content-Type": "application/json"}

def fetch_rest(endpoint: str):
    url = f"{SUPABASE_URL}/rest/v1/{endpoint}"
    req = urllib.request.Request(url, headers=headers)
    with urllib.request.urlopen(req, context=ctx) as r:
        return json.loads(r.read().decode("utf-8"))

async def generate_comprehensive_report():
    print("Fetching data from Supabase...")
    units = fetch_rest("micro_units?select=*&order=unit_number.asc")
    creators = fetch_rest("micro_unit_creators?select=*&order=name.asc")
    channels = fetch_rest("micro_unit_channels?select=*")
    ig_metrics = []
    for ym in ["2026-08", "2026-09"]:
        ig_metrics.extend(fetch_rest(f"monthly_channel_metrics?select=*&year_month=eq.{ym}"))
    users = fetch_rest("users?select=id,full_name,email")

    user_map = {u["id"]: u["full_name"] for u in users}

    # Map Instagram metrics
    ig_map = {}
    for m in ig_metrics:
        ym = m["year_month"]
        ig_id = str(m.get("instagram_id") or "").strip()
        u = str(m.get("username") or "").strip().lower().lstrip("@")
        if ig_id:
            ig_map.setdefault(ig_id, {})[ym] = m
            if ig_id.startswith("sp_"):
                ig_map.setdefault(ig_id[3:], {})[ym] = m
        if u:
            ig_map.setdefault(u, {})[ym] = m

    yt_channel_ids = [
        c["username"] for c in channels 
        if (c.get("platform") or "INSTAGRAM").upper() == "YOUTUBE"
    ]
    print(f"Calculating YouTube monthly metrics for {len(yt_channel_ids)} channels...")
    yt_aug = await calculate_youtube_monthly_metrics(yt_channel_ids, 2026, 8)
    yt_sep = await calculate_youtube_monthly_metrics(yt_channel_ids, 2026, 9)

    unit_creators_map = defaultdict(list)
    for cr in creators:
        unit_creators_map[cr["micro_unit_id"]].append(cr["name"])

    unit_channels_map = defaultdict(list)
    for ch in channels:
        unit_channels_map[ch["micro_unit_id"]].append(ch)

    processed_units = []
    grand_aug_views = 0.0
    grand_sep_views = 0.0
    grand_aug_posts = 0
    grand_sep_posts = 0

    grand_ig_aug_v = 0.0
    grand_ig_aug_p = 0
    grand_ig_sep_v = 0.0
    grand_ig_sep_p = 0

    grand_yt_aug_v = 0.0
    grand_yt_aug_p = 0
    grand_yt_sep_v = 0.0
    grand_yt_sep_p = 0

    for u in units:
        u_id = u["id"]
        poc_name = user_map.get(u.get("poc_user_id")) or "Unassigned"
        unit_chs = unit_channels_map.get(u_id, [])
        explicit_cr_names = unit_creators_map.get(u_id, [])

        creator_groups = defaultdict(list)
        for c_name in explicit_cr_names:
            creator_groups[c_name] = []

        for ch in unit_chs:
            c_name = (ch.get("creator_name") or ch.get("channel_title") or ch.get("username") or "Unassigned Creator").strip()
            creator_groups[c_name].append(ch)

        # Sort creators: POC FIRST, then alphabetical
        def creator_sort_key(name):
            is_poc = poc_name and name.strip().lower() == poc_name.strip().lower()
            return (0 if is_poc else 1, name.lower())

        sorted_creator_names = sorted(creator_groups.keys(), key=creator_sort_key)

        creators_list = []
        unit_aug_views = 0.0
        unit_sep_views = 0.0
        unit_aug_posts = 0
        unit_sep_posts = 0

        unit_ig_aug_v, unit_ig_aug_p = 0.0, 0
        unit_ig_sep_v, unit_ig_sep_p = 0.0, 0
        unit_yt_aug_v, unit_yt_aug_p = 0.0, 0
        unit_yt_sep_v, unit_yt_sep_p = 0.0, 0

        for cr_name in sorted_creator_names:
            ch_list = creator_groups[cr_name]
            is_poc = bool(poc_name and cr_name.strip().lower() == poc_name.strip().lower())
            
            cr_ig_aug_v, cr_ig_aug_p = 0.0, 0
            cr_ig_sep_v, cr_ig_sep_p = 0.0, 0
            cr_yt_aug_v, cr_yt_aug_p = 0.0, 0
            cr_yt_sep_v, cr_yt_sep_p = 0.0, 0

            yt_channel_names = []
            ig_channel_names = []

            for ch in ch_list:
                platform = (ch.get("platform") or "INSTAGRAM").upper()
                raw_handle = ch.get("username") or ""
                clean_handle = raw_handle.lstrip("@").strip()
                title = ch.get("channel_title") or raw_handle

                if platform == "YOUTUBE":
                    m_aug = yt_aug.get(raw_handle, {})
                    m_sep = yt_sep.get(raw_handle, {})
                    aug_v = float(m_aug.get("views", 0.0) or 0.0)
                    sep_v = float(m_sep.get("views", 0.0) or 0.0)
                    aug_p = int(m_aug.get("video_count", 0) or 0)
                    sep_p = int(m_sep.get("video_count", 0) or 0)
                    title = m_aug.get("title") or title
                    yt_channel_names.append(title)
                    cr_yt_aug_v += aug_v
                    cr_yt_aug_p += aug_p
                    cr_yt_sep_v += sep_v
                    cr_yt_sep_p += sep_p
                else:
                    ig_id = str(ch.get("instagram_id") or "").strip()
                    m_aug = ig_map.get(ig_id, {}).get("2026-08")
                    if not m_aug and ig_id.startswith("sp_"):
                        m_aug = ig_map.get(ig_id[3:], {}).get("2026-08")
                    if not m_aug:
                        m_aug = ig_map.get(clean_handle.lower(), {}).get("2026-08")

                    m_sep = ig_map.get(ig_id, {}).get("2026-09")
                    if not m_sep and ig_id.startswith("sp_"):
                        m_sep = ig_map.get(ig_id[3:], {}).get("2026-09")
                    if not m_sep:
                        m_sep = ig_map.get(clean_handle.lower(), {}).get("2026-09")

                    aug_v = float(m_aug.get("monthly_views") or 0.0) if m_aug else 0.0
                    sep_v = float(m_sep.get("monthly_views") or 0.0) if m_sep else 0.0
                    aug_p = int(m_aug.get("reels_count") if (m_aug and m_aug.get("reels_count") is not None) else (m_aug.get("post_count", 0) if m_aug else 0))
                    sep_p = int(m_sep.get("reels_count") if (m_sep and m_sep.get("reels_count") is not None) else (m_sep.get("post_count", 0) if m_sep else 0))
                    ig_channel_names.append(f"@{clean_handle}")
                    cr_ig_aug_v += aug_v
                    cr_ig_aug_p += aug_p
                    cr_ig_sep_v += sep_v
                    cr_ig_sep_p += sep_p

            tot_cr_aug_v = cr_ig_aug_v + cr_yt_aug_v
            tot_cr_sep_v = cr_ig_sep_v + cr_yt_sep_v
            tot_cr_aug_p = cr_ig_aug_p + cr_yt_aug_p
            tot_cr_sep_p = cr_ig_sep_p + cr_yt_sep_p

            unit_aug_views += tot_cr_aug_v
            unit_sep_views += tot_cr_sep_v
            unit_aug_posts += tot_cr_aug_p
            unit_sep_posts += tot_cr_sep_p

            unit_ig_aug_v += cr_ig_aug_v
            unit_ig_aug_p += cr_ig_aug_p
            unit_ig_sep_v += cr_ig_sep_v
            unit_ig_sep_p += cr_ig_sep_p

            unit_yt_aug_v += cr_yt_aug_v
            unit_yt_aug_p += cr_yt_aug_p
            unit_yt_sep_v += cr_yt_sep_v
            unit_yt_sep_p += cr_yt_sep_p

            creators_list.append({
                "creator_name": cr_name,
                "is_poc": is_poc,
                "channels_count": len(ch_list),
                "yt_channels_str": ", ".join(yt_channel_names) if yt_channel_names else "—",
                "ig_channels_str": ", ".join(ig_channel_names) if ig_channel_names else "—",
                "ig_aug_v": cr_ig_aug_v,
                "ig_aug_p": cr_ig_aug_p,
                "ig_sep_v": cr_ig_sep_v,
                "ig_sep_p": cr_ig_sep_p,
                "yt_aug_v": cr_yt_aug_v,
                "yt_aug_p": cr_yt_aug_p,
                "yt_sep_v": cr_yt_sep_v,
                "yt_sep_p": cr_yt_sep_p,
                "aug_views": tot_cr_aug_v,
                "aug_posts": tot_cr_aug_p,
                "sep_views": tot_cr_sep_v,
                "sep_posts": tot_cr_sep_p,
                "total_views": tot_cr_aug_v + tot_cr_sep_v,
                "total_posts": tot_cr_aug_p + tot_cr_sep_p
            })

        # Calculate unit-level winners
        def get_max_cr(key):
            filtered = [c for c in creators_list if c[key] > 0]
            if not filtered:
                return None
            return max(filtered, key=lambda c: c[key])

        highlights = {
            "ig_aug_views": get_max_cr("ig_aug_v"),
            "ig_aug_reels": get_max_cr("ig_aug_p"),
            "ig_sep_views": get_max_cr("ig_sep_v"),
            "ig_sep_reels": get_max_cr("ig_sep_p"),
            "yt_aug_views": get_max_cr("yt_aug_v"),
            "yt_aug_videos": get_max_cr("yt_aug_p"),
            "yt_sep_views": get_max_cr("yt_sep_v"),
            "yt_sep_videos": get_max_cr("yt_sep_p"),
        }

        grand_aug_views += unit_aug_views
        grand_sep_views += unit_sep_views
        grand_aug_posts += unit_aug_posts
        grand_sep_posts += unit_sep_posts

        grand_ig_aug_v += unit_ig_aug_v
        grand_ig_aug_p += unit_ig_aug_p
        grand_ig_sep_v += unit_ig_sep_v
        grand_ig_sep_p += unit_ig_sep_p

        grand_yt_aug_v += unit_yt_aug_v
        grand_yt_aug_p += unit_yt_aug_p
        grand_yt_sep_v += unit_yt_sep_v
        grand_yt_sep_p += unit_yt_sep_p

        processed_units.append({
            "unit_number": u.get("unit_number"),
            "name": u.get("name"),
            "poc_name": poc_name,
            "creators": creators_list,
            "highlights": highlights,
            "total_channels": len(unit_chs),
            "aug_views": unit_aug_views,
            "aug_posts": unit_aug_posts,
            "sep_views": unit_sep_views,
            "sep_posts": unit_sep_posts,
            "total_views": unit_aug_views + unit_sep_views,
            "total_posts": unit_aug_posts + unit_sep_posts,
            "ig_aug_v": unit_ig_aug_v,
            "ig_aug_p": unit_ig_aug_p,
            "ig_sep_v": unit_ig_sep_v,
            "ig_sep_p": unit_ig_sep_p,
            "yt_aug_v": unit_yt_aug_v,
            "yt_aug_p": unit_yt_aug_p,
            "yt_sep_v": unit_yt_sep_v,
            "yt_sep_p": unit_yt_sep_p,
        })

    # Identify Overall Top Performing Units
    top_aug_unit = max(processed_units, key=lambda x: x["aug_views"])
    top_sep_unit = max(processed_units, key=lambda x: x["sep_views"])

    # ---------------- BUILD EXCEL WORKBOOK ----------------
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Micro Units Performance"
    ws.views.sheetView[0].showGridLines = True

    # Styling Palette
    c_banner_bg = "1E1B4B"       # Deep Purple-Navy
    c_tbl_header = "4C1D95"      # Deep Purple Header
    c_ig_header = "831843"       # Rich Rose/Wine for Instagram
    c_yt_header = "881337"       # Crimson for YouTube
    c_tot_header = "1E1B4B"      # Dark Navy for Totals
    c_sec_header = "31104B"      # Dark Violet Section
    c_zebra = "F8FAFC"           # Light Slate
    c_unit_header_bg = "EDE9FE"  # Soft Lavender for Unit Title
    c_poc_bg = "FAF5FF"          # Light Purple for POC Row
    c_subtotal_bg = "F1F5F9"     # Light Slate Subtotal
    c_grand_total_bg = "FDE68A"  # Amber 200

    # Highlight Colors
    c_gold_highlight = "FEF08A"   # Bright Gold for Max Views
    c_reel_highlight = "DDD6FE"   # Lavender Violet for Max Reels
    c_yt_video_hl = "FED7AA"      # Warm Peach for Max YT Videos
    c_top_unit_aug = "FEF08A"     # Gold for August Max Unit
    c_top_unit_sep = "BBF7D0"     # Fresh Emerald for September Max Unit

    font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="E0E7FF")
    font_kpi_label = Font(name="Calibri", size=9, bold=True, color="4B5563")
    font_kpi_val = Font(name="Calibri", size=14, bold=True, color="1E1B4B")
    font_sec_header = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    font_tbl_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_sub_header = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
    
    font_data = Font(name="Calibri", size=10)
    font_data_poc = Font(name="Calibri", size=10, bold=True, color="4C1D95")
    font_data_bold = Font(name="Calibri", size=10, bold=True)
    font_unit_header = Font(name="Calibri", size=11, bold=True, color="2E1065")
    font_unit_callout = Font(name="Calibri", size=9, bold=True, color="3730A3")
    font_unit_subtotal = Font(name="Calibri", size=10, bold=True, color="31104B")
    font_grand_total = Font(name="Calibri", size=11, bold=True, color="78350F")

    font_winner_views = Font(name="Calibri", size=10, bold=True, color="854D0E")
    font_winner_reels = Font(name="Calibri", size=10, bold=True, color="5B21B6")
    font_winner_videos = Font(name="Calibri", size=10, bold=True, color="9A3412")

    fill_banner = PatternFill(start_color=c_banner_bg, end_color=c_banner_bg, fill_type="solid")
    fill_tbl_header = PatternFill(start_color=c_tbl_header, end_color=c_tbl_header, fill_type="solid")
    fill_ig_header = PatternFill(start_color=c_ig_header, end_color=c_ig_header, fill_type="solid")
    fill_yt_header = PatternFill(start_color=c_yt_header, end_color=c_yt_header, fill_type="solid")
    fill_tot_header = PatternFill(start_color=c_tot_header, end_color=c_tot_header, fill_type="solid")
    fill_sec_header = PatternFill(start_color=c_sec_header, end_color=c_sec_header, fill_type="solid")
    fill_zebra = PatternFill(start_color=c_zebra, end_color=c_zebra, fill_type="solid")
    fill_unit_header = PatternFill(start_color=c_unit_header_bg, end_color=c_unit_header_bg, fill_type="solid")
    fill_unit_callout = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
    fill_poc_row = PatternFill(start_color=c_poc_bg, end_color=c_poc_bg, fill_type="solid")
    fill_unit_subtotal = PatternFill(start_color=c_subtotal_bg, end_color=c_subtotal_bg, fill_type="solid")
    fill_grand_total = PatternFill(start_color=c_grand_total_bg, end_color=c_grand_total_bg, fill_type="solid")

    fill_hl_views = PatternFill(start_color=c_gold_highlight, end_color=c_gold_highlight, fill_type="solid")
    fill_hl_reels = PatternFill(start_color=c_reel_highlight, end_color=c_reel_highlight, fill_type="solid")
    fill_hl_videos = PatternFill(start_color=c_yt_video_hl, end_color=c_yt_video_hl, fill_type="solid")
    fill_top_aug = PatternFill(start_color=c_top_unit_aug, end_color=c_top_unit_aug, fill_type="solid")
    fill_top_sep = PatternFill(start_color=c_top_unit_sep, end_color=c_top_unit_sep, fill_type="solid")

    thin_border = Side(style='thin', color='D1D5DB')
    thick_top = Side(style='medium', color='4C1D95')
    double_bottom = Side(style='double', color='1E1B4B')

    border_cell = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    border_grand_total = Border(left=thin_border, right=thin_border, top=thick_top, bottom=double_bottom)

    # STRICT CENTER ALIGNMENT FOR ALL CELLS
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    num_fmt_int = "#,##0"
    num_fmt_views = "#,##0"
    num_fmt_pct = "0.0%"

    TOTAL_COLS = 16
    row_idx = 1

    # 1. TITLE BANNER
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx+1, end_column=TOTAL_COLS)
    b_cell = ws.cell(row=row_idx, column=1, value="Sadhguru Wisdom Warriors — Micro Units Multi-Platform Performance Report")
    b_cell.font = font_title
    b_cell.alignment = align_center
    for r in range(row_idx, row_idx+2):
        for c in range(1, TOTAL_COLS+1):
            ws.cell(row=r, column=c).fill = fill_banner
    row_idx += 2

    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=TOTAL_COLS)
    sub_cell = ws.cell(row=row_idx, column=1, value=f"August 2026 & September 2026 Reach Across All 11 Micro Units (Instagram + YouTube) | Generated: {datetime.now().strftime('%B %d, %Y')}")
    sub_cell.font = font_subtitle
    sub_cell.alignment = align_center
    for c in range(1, TOTAL_COLS+1):
        ws.cell(row=row_idx, column=c).fill = fill_banner
    row_idx += 2

    # 2. EXECUTIVE KPI CARDS
    kpis = [
        ("TOTAL UNITS", len(processed_units), num_fmt_int),
        ("TOTAL CREATORS", sum(len(u["creators"]) for u in processed_units), num_fmt_int),
        ("TOTAL CHANNELS", sum(u["total_channels"] for u in processed_units), num_fmt_int),
        ("AUG 2026 TOTAL VIEWS", grand_aug_views, num_fmt_views),
        ("SEP 2026 TOTAL VIEWS", grand_sep_views, num_fmt_views),
        ("COMBINED TOTAL VIEWS", grand_aug_views + grand_sep_views, num_fmt_views),
    ]

    col_width_per_kpi = TOTAL_COLS // len(kpis)
    col_start = 1
    for i, (label, val, fmt) in enumerate(kpis):
        c1 = col_start
        c2 = col_start + col_width_per_kpi - 1 if i < len(kpis) - 1 else TOTAL_COLS
        ws.merge_cells(start_row=row_idx, start_column=c1, end_row=row_idx, end_column=c2)
        ws.merge_cells(start_row=row_idx+1, start_column=c1, end_row=row_idx+1, end_column=c2)

        l_cell = ws.cell(row=row_idx, column=c1, value=label)
        l_cell.font = font_kpi_label
        l_cell.alignment = align_center
        l_cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        v_cell = ws.cell(row=row_idx+1, column=c1, value=val)
        v_cell.font = font_kpi_val
        v_cell.number_format = fmt
        v_cell.alignment = align_center
        v_cell.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        for r in range(row_idx, row_idx+2):
            for c in range(c1, c2+1):
                ws.cell(row=r, column=c).border = border_cell

        col_start = c2 + 1

    row_idx += 3

    # 2B. DEDICATED PLATFORM TOTALS SUMMARY TABLE (YOUTUBE & INSTAGRAM SEPARATELY)
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=TOTAL_COLS)
    p_sec_title = ws.cell(row=row_idx, column=1, value="🌐 PLATFORM PERFORMANCE TOTALS — INSTAGRAM & YOUTUBE SEPARATELY")
    p_sec_title.font = font_sec_header
    p_sec_title.fill = fill_sec_header
    p_sec_title.alignment = align_center
    for c in range(1, TOTAL_COLS+1):
        ws.cell(row=row_idx, column=c).fill = fill_sec_header
    row_idx += 1

    plat_headers = [
        (1, 3, "Platform Name"),
        (4, 5, "Aug 2026 Views"),
        (6, 7, "Aug Content Items"),
        (8, 9, "Sep 2026 Views"),
        (10, 11, "Sep Content Items"),
        (12, 13, "Total Views (Aug+Sep)"),
        (14, 15, "Total Content Items"),
        (16, 16, "% Share of Reach"),
    ]
    for c1, c2, h_text in plat_headers:
        if c1 != c2:
            ws.merge_cells(start_row=row_idx, start_column=c1, end_row=row_idx, end_column=c2)
        cell = ws.cell(row=row_idx, column=c1, value=h_text)
        cell.font = font_tbl_header
        cell.fill = fill_tbl_header
        cell.alignment = align_center
        for c in range(c1, c2+1):
            ws.cell(row=row_idx, column=c).fill = fill_tbl_header
            ws.cell(row=row_idx, column=c).border = border_cell
    row_idx += 1

    grand_tot_v = grand_aug_views + grand_sep_views
    ig_share = (grand_ig_aug_v + grand_ig_sep_v) / grand_tot_v if grand_tot_v > 0 else 0
    yt_share = (grand_yt_aug_v + grand_yt_sep_v) / grand_tot_v if grand_tot_v > 0 else 0

    plat_rows = [
        ("📸 Instagram", grand_ig_aug_v, grand_ig_aug_p, grand_ig_sep_v, grand_ig_sep_p, grand_ig_aug_v + grand_ig_sep_v, grand_ig_aug_p + grand_ig_sep_p, ig_share, fill_ig_header, PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid")),
        ("▶️ YouTube", grand_yt_aug_v, grand_yt_aug_p, grand_yt_sep_v, grand_yt_sep_p, grand_yt_aug_v + grand_yt_sep_v, grand_yt_aug_p + grand_yt_sep_p, yt_share, fill_yt_header, PatternFill(start_color="FFF7ED", end_color="FFF7ED", fill_type="solid")),
        ("🌐 Combined Total", grand_aug_views, grand_aug_posts, grand_sep_views, grand_sep_posts, grand_tot_v, grand_aug_posts + grand_sep_posts, 1.0, fill_tot_header, fill_grand_total),
    ]

    for p_name, a_v, a_p, s_v, s_p, t_v, t_p, sh, hdr_fill, row_fill in plat_rows:
        is_tot = "Combined" in p_name
        p_font = font_grand_total if is_tot else font_data_bold
        p_row_data = [
            (1, 3, p_name, "@"),
            (4, 5, a_v, num_fmt_views),
            (6, 7, a_p, num_fmt_int),
            (8, 9, s_v, num_fmt_views),
            (10, 11, s_p, num_fmt_int),
            (12, 13, t_v, num_fmt_views),
            (14, 15, t_p, num_fmt_int),
            (16, 16, sh, num_fmt_pct),
        ]
        for c1, c2, val, fmt in p_row_data:
            if c1 != c2:
                ws.merge_cells(start_row=row_idx, start_column=c1, end_row=row_idx, end_column=c2)
            cell = ws.cell(row=row_idx, column=c1, value=val)
            cell.font = p_font
            cell.number_format = fmt
            cell.alignment = align_center
            for c in range(c1, c2+1):
                ws.cell(row=row_idx, column=c).fill = row_fill
                ws.cell(row=row_idx, column=c).border = border_grand_total if is_tot else border_cell
        row_idx += 1

    row_idx += 2

    # 3. EXECUTIVE SUMMARY TABLE (With August & September Max Units Highlighted and Platform Columns)
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=TOTAL_COLS)
    s_title = ws.cell(row=row_idx, column=1, value="1. EXECUTIVE SUMMARY — ALL 11 MICRO UNITS MATRIX (PLATFORMS SEPARATED)")
    s_title.font = font_sec_header
    s_title.fill = fill_sec_header
    s_title.alignment = align_center
    for c in range(1, TOTAL_COLS+1):
        ws.cell(row=row_idx, column=c).fill = fill_sec_header
    row_idx += 1

    # Callout banner for Top Units
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=TOTAL_COLS)
    c_banner = ws.cell(
        row=row_idx, column=1, 
        value=f"🏆 HIGHLIGHTS: [August Top Unit: {top_aug_unit['name']} ({top_aug_unit['aug_views']:,.0f} views)]  |  [September Top Unit: {top_sep_unit['name']} ({top_sep_unit['sep_views']:,.0f} views)]"
    )
    c_banner.font = Font(name="Calibri", size=10, bold=True, color="1E1B4B")
    c_banner.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    c_banner.alignment = align_center
    for c in range(1, TOTAL_COLS+1):
        ws.cell(row=row_idx, column=c).border = border_cell
        ws.cell(row=row_idx, column=c).fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    row_idx += 1

    summary_headers = [
        "Unit #", "Micro Unit Name", "Assigned POC", "Creators", "Channels", 
        "📸 Aug IG Views", "▶️ Aug YT Views", "Aug Total Views", 
        "📸 Sep IG Views", "▶️ Sep YT Views", "Sep Total Views", 
        "Total IG Views", "Total YT Views", "Combined Views", "Total Content", "% Reach"
    ]
    for col_i, h in enumerate(summary_headers, start=1):
        cell = ws.cell(row=row_idx, column=col_i, value=h)
        cell.font = font_tbl_header
        cell.alignment = align_center
        cell.border = border_cell
        if col_i in (6, 9, 12): cell.fill = fill_ig_header
        elif col_i in (7, 10, 13): cell.fill = fill_yt_header
        else: cell.fill = fill_tbl_header
    row_idx += 1

    sum_start_row = row_idx
    combined_grand_views = grand_aug_views + grand_sep_views

    for idx, u in enumerate(processed_units):
        pct = (u["total_views"] / combined_grand_views) if combined_grand_views > 0 else 0
        is_top_aug = (u["name"] == top_aug_unit["name"])
        is_top_sep = (u["name"] == top_sep_unit["name"])

        tot_u_ig_v = u["ig_aug_v"] + u["ig_sep_v"]
        tot_u_yt_v = u["yt_aug_v"] + u["yt_sep_v"]

        vals = [
            (u["unit_number"], num_fmt_int),
            (u["name"], "@"),
            (u["poc_name"], "@"),
            (len(u["creators"]), num_fmt_int),
            (u["total_channels"], num_fmt_int),
            (u["ig_aug_v"], num_fmt_views),
            (u["yt_aug_v"], num_fmt_views),
            (u["aug_views"], num_fmt_views),
            (u["ig_sep_v"], num_fmt_views),
            (u["yt_sep_v"], num_fmt_views),
            (u["sep_views"], num_fmt_views),
            (tot_u_ig_v, num_fmt_views),
            (tot_u_yt_v, num_fmt_views),
            (u["total_views"], num_fmt_views),
            (u["total_posts"], num_fmt_int),
            (pct, num_fmt_pct),
        ]

        fill_row = fill_zebra if idx % 2 == 1 else PatternFill(fill_type=None)

        for col_i, (v, fmt) in enumerate(vals, start=1):
            cell = ws.cell(row=row_idx, column=col_i, value=v)
            cell.font = font_data
            cell.number_format = fmt
            cell.alignment = align_center
            cell.border = border_cell

            # Default zebra
            if fill_row.fill_type:
                cell.fill = fill_row

            # HIGHLIGHT TOP AUGUST MICRO UNIT
            if is_top_aug and col_i in (2, 8):
                cell.fill = fill_top_aug
                cell.font = font_winner_views
            
            # HIGHLIGHT TOP SEPTEMBER MICRO UNIT
            if is_top_sep and col_i in (2, 11):
                cell.fill = fill_top_sep
                cell.font = Font(name="Calibri", size=10, bold=True, color="065F46")

        row_idx += 1

    # Summary Grand Total Row
    sum_tot_vals = [
        ("TOTAL", "@"),
        ("All 11 Micro Units", "@"),
        ("—", "@"),
        (f"=SUM(D{sum_start_row}:D{row_idx-1})", num_fmt_int),
        (f"=SUM(E{sum_start_row}:E{row_idx-1})", num_fmt_int),
        (f"=SUM(F{sum_start_row}:F{row_idx-1})", num_fmt_views),  # Aug IG Views Total
        (f"=SUM(G{sum_start_row}:G{row_idx-1})", num_fmt_views),  # Aug YT Views Total
        (f"=SUM(H{sum_start_row}:H{row_idx-1})", num_fmt_views),  # Aug Total Views
        (f"=SUM(I{sum_start_row}:I{row_idx-1})", num_fmt_views),  # Sep IG Views Total
        (f"=SUM(J{sum_start_row}:J{row_idx-1})", num_fmt_views),  # Sep YT Views Total
        (f"=SUM(K{sum_start_row}:K{row_idx-1})", num_fmt_views),  # Sep Total Views
        (f"=SUM(L{sum_start_row}:L{row_idx-1})", num_fmt_views),  # Total IG Views
        (f"=SUM(M{sum_start_row}:M{row_idx-1})", num_fmt_views),  # Total YT Views
        (f"=SUM(N{sum_start_row}:N{row_idx-1})", num_fmt_views),  # Combined Total Views
        (f"=SUM(O{sum_start_row}:O{row_idx-1})", num_fmt_int),    # Total Content
        (1.0, num_fmt_pct),
    ]

    for col_i, (v, fmt) in enumerate(sum_tot_vals, start=1):
        cell = ws.cell(row=row_idx, column=col_i, value=v)
        cell.font = font_grand_total
        cell.number_format = fmt
        cell.alignment = align_center
        cell.fill = fill_grand_total
        cell.border = border_grand_total

    row_idx += 3

    # 4. DETAILED BREAKDOWN TABLE — SEPARATING INSTAGRAM & YOUTUBE WITH PLATFORM WINNERS HIGHLIGHTED
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=TOTAL_COLS)
    d_title = ws.cell(row=row_idx, column=1, value="2. DETAILED BREAKDOWN — PLATFORM PERFORMANCE (INSTAGRAM & YOUTUBE) BY MICRO UNIT")
    d_title.font = font_sec_header
    d_title.fill = fill_sec_header
    d_title.alignment = align_center
    for c in range(1, TOTAL_COLS+1):
        ws.cell(row=row_idx, column=c).fill = fill_sec_header
    row_idx += 1

    # Helper function to render table headers for each micro unit
    def render_unit_table_headers(r_idx):
        # Tier 1: Group Headers
        ws.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=4)
        c_meta = ws.cell(row=r_idx, column=1, value="CREATOR PROFILE")
        c_meta.font = font_tbl_header
        c_meta.fill = fill_tbl_header
        c_meta.alignment = align_center
        for c in range(1, 5): 
            ws.cell(row=r_idx, column=c).fill = fill_tbl_header
            ws.cell(row=r_idx, column=c).border = border_cell

        ws.merge_cells(start_row=r_idx, start_column=5, end_row=r_idx, end_column=8)
        c_ig = ws.cell(row=r_idx, column=5, value="📸 INSTAGRAM PERFORMANCE (REELS & VIEWS)")
        c_ig.font = font_tbl_header
        c_ig.fill = fill_ig_header
        c_ig.alignment = align_center
        for c in range(5, 9): 
            ws.cell(row=r_idx, column=c).fill = fill_ig_header
            ws.cell(row=r_idx, column=c).border = border_cell

        ws.merge_cells(start_row=r_idx, start_column=9, end_row=r_idx, end_column=12)
        c_yt = ws.cell(row=r_idx, column=9, value="▶️ YOUTUBE PERFORMANCE (VIDEOS & VIEWS)")
        c_yt.font = font_tbl_header
        c_yt.fill = fill_yt_header
        c_yt.alignment = align_center
        for c in range(9, 13): 
            ws.cell(row=r_idx, column=c).fill = fill_yt_header
            ws.cell(row=r_idx, column=c).border = border_cell

        ws.merge_cells(start_row=r_idx, start_column=13, end_row=r_idx, end_column=16)
        c_all = ws.cell(row=r_idx, column=13, value="🌐 COMBINED OVERALL TOTALS")
        c_all.font = font_tbl_header
        c_all.fill = fill_tot_header
        c_all.alignment = align_center
        for c in range(13, 17): 
            ws.cell(row=r_idx, column=c).fill = fill_tot_header
            ws.cell(row=r_idx, column=c).border = border_cell

        r_idx += 1

        # Tier 2: Column Headers
        detail_headers = [
            "Micro Unit", "Role / Status", "Wisdom Warrior Name", "Channels",
            "Aug IG Views", "Aug IG Reels", "Sep IG Views", "Sep IG Reels",
            "Aug YT Views", "Aug YT Videos", "Sep YT Views", "Sep YT Videos",
            "Aug Total Views", "Sep Total Views", "Combined Views", "Total Content"
        ]
        for col_i, h in enumerate(detail_headers, start=1):
            cell = ws.cell(row=r_idx, column=col_i, value=h)
            cell.font = font_sub_header
            cell.alignment = align_center
            cell.border = border_cell
            if col_i in range(1, 5): cell.fill = fill_tbl_header
            elif col_i in range(5, 9): cell.fill = fill_ig_header
            elif col_i in range(9, 13): cell.fill = fill_yt_header
            else: cell.fill = fill_tot_header

        r_idx += 1
        return r_idx

    row_idx += 1

    # Iterate Through Units
    for u in processed_units:
        u_name = u["name"]
        poc = u["poc_name"]
        hl = u["highlights"]

        # Unit Section Header
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=TOTAL_COLS)
        u_sec = ws.cell(row=row_idx, column=1, value=f"⭐ {u_name} — POC: {poc}  |  Creators: {len(u['creators'])}, Channels: {u['total_channels']}")
        u_sec.font = font_unit_header
        u_sec.fill = fill_unit_header
        u_sec.alignment = align_center
        for c in range(1, TOTAL_COLS+1):
            ws.cell(row=row_idx, column=c).fill = fill_unit_header
            ws.cell(row=row_idx, column=c).border = border_cell
        row_idx += 1

        # Render Table Header specifically for this Micro Unit
        row_idx = render_unit_table_headers(row_idx)

        # Creators in Unit
        for cr in u["creators"]:
            cr_name = cr["creator_name"]
            is_poc = cr["is_poc"]
            role_label = "POC" if is_poc else "Team Member"
            display_name = f"{cr_name} [POC]" if is_poc else cr_name

            # Check if this creator is a platform winner in this unit
            win_ig_aug_v = bool(hl["ig_aug_views"] and hl["ig_aug_views"]["creator_name"] == cr_name and cr["ig_aug_v"] > 0)
            win_ig_aug_p = bool(hl["ig_aug_reels"] and hl["ig_aug_reels"]["creator_name"] == cr_name and cr["ig_aug_p"] > 0)
            win_ig_sep_v = bool(hl["ig_sep_views"] and hl["ig_sep_views"]["creator_name"] == cr_name and cr["ig_sep_v"] > 0)
            win_ig_sep_p = bool(hl["ig_sep_reels"] and hl["ig_sep_reels"]["creator_name"] == cr_name and cr["ig_sep_p"] > 0)

            win_yt_aug_v = bool(hl["yt_aug_views"] and hl["yt_aug_views"]["creator_name"] == cr_name and cr["yt_aug_v"] > 0)
            win_yt_aug_p = bool(hl["yt_aug_videos"] and hl["yt_aug_videos"]["creator_name"] == cr_name and cr["yt_aug_p"] > 0)
            win_yt_sep_v = bool(hl["yt_sep_views"] and hl["yt_sep_views"]["creator_name"] == cr_name and cr["yt_sep_v"] > 0)
            win_yt_sep_p = bool(hl["yt_sep_videos"] and hl["yt_sep_videos"]["creator_name"] == cr_name and cr["yt_sep_p"] > 0)

            row_data = [
                (u_name, "@"),
                (role_label, "@"),
                (display_name, "@"),
                (cr["channels_count"], num_fmt_int),

                # Instagram
                (cr["ig_aug_v"], num_fmt_views),
                (cr["ig_aug_p"], num_fmt_int),
                (cr["ig_sep_v"], num_fmt_views),
                (cr["ig_sep_p"], num_fmt_int),

                # YouTube
                (cr["yt_aug_v"], num_fmt_views),
                (cr["yt_aug_p"], num_fmt_int),
                (cr["yt_sep_v"], num_fmt_views),
                (cr["yt_sep_p"], num_fmt_int),

                # Combined
                (cr["aug_views"], num_fmt_views),
                (cr["sep_views"], num_fmt_views),
                (cr["total_views"], num_fmt_views),
                (cr["total_posts"], num_fmt_int),
            ]

            fill_cr = fill_poc_row if is_poc else PatternFill(fill_type=None)
            font_cr = font_data_poc if is_poc else font_data

            for col_i, (v, fmt) in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_i, value=v)
                cell.font = font_cr
                cell.number_format = fmt
                cell.alignment = align_center
                cell.border = border_cell
                if fill_cr.fill_type:
                    cell.fill = fill_cr

                # APPLY SEPARATE PLATFORM HIGHLIGHTS AS REQUESTED
                # IG Aug Max Views
                if col_i == 5 and win_ig_aug_v:
                    cell.fill = fill_hl_views
                    cell.font = font_winner_views
                # IG Aug Max Reels
                elif col_i == 6 and win_ig_aug_p:
                    cell.fill = fill_hl_reels
                    cell.font = font_winner_reels
                # IG Sep Max Views
                elif col_i == 7 and win_ig_sep_v:
                    cell.fill = fill_hl_views
                    cell.font = font_winner_views
                # IG Sep Max Reels
                elif col_i == 8 and win_ig_sep_p:
                    cell.fill = fill_hl_reels
                    cell.font = font_winner_reels

                # YT Aug Max Views
                elif col_i == 9 and win_yt_aug_v:
                    cell.fill = fill_hl_views
                    cell.font = font_winner_views
                # YT Aug Max Videos
                elif col_i == 10 and win_yt_aug_p:
                    cell.fill = fill_hl_videos
                    cell.font = font_winner_videos
                # YT Sep Max Views
                elif col_i == 11 and win_yt_sep_v:
                    cell.fill = fill_hl_views
                    cell.font = font_winner_views
                # YT Sep Max Videos
                elif col_i == 12 and win_yt_sep_p:
                    cell.fill = fill_hl_videos
                    cell.font = font_winner_videos

            row_idx += 1

        # Unit Subtotal Row
        sub_vals = [
            (u_name, "@"),
            ("SUBTOTAL", "@"),
            (f"{u_name} Total ({len(u['creators'])} Members)", "@"),
            (u["total_channels"], num_fmt_int),
            (u["ig_aug_v"], num_fmt_views),
            (u["ig_aug_p"], num_fmt_int),
            (u["ig_sep_v"], num_fmt_views),
            (u["ig_sep_p"], num_fmt_int),
            (u["yt_aug_v"], num_fmt_views),
            (u["yt_aug_p"], num_fmt_int),
            (u["yt_sep_v"], num_fmt_views),
            (u["yt_sep_p"], num_fmt_int),
            (u["aug_views"], num_fmt_views),
            (u["sep_views"], num_fmt_views),
            (u["total_views"], num_fmt_views),
            (u["total_posts"], num_fmt_int),
        ]

        for col_i, (v, fmt) in enumerate(sub_vals, start=1):
            cell = ws.cell(row=row_idx, column=col_i, value=v)
            cell.font = font_unit_subtotal
            cell.number_format = fmt
            cell.alignment = align_center
            cell.fill = fill_unit_subtotal
            cell.border = border_cell

        row_idx += 2

    # Grand Total Row for Detailed Breakdown
    final_grand_vals = [
        ("GRAND TOTAL", "@"),
        ("All Micro Units", "@"),
        (f"54 Creators Across 11 Units", "@"),
        (sum(u["total_channels"] for u in processed_units), num_fmt_int),
        (grand_ig_aug_v, num_fmt_views),
        (grand_ig_aug_p, num_fmt_int),
        (grand_ig_sep_v, num_fmt_views),
        (grand_ig_sep_p, num_fmt_int),
        (grand_yt_aug_v, num_fmt_views),
        (grand_yt_aug_p, num_fmt_int),
        (grand_yt_sep_v, num_fmt_views),
        (grand_yt_sep_p, num_fmt_int),
        (grand_aug_views, num_fmt_views),
        (grand_sep_views, num_fmt_views),
        (grand_aug_views + grand_sep_views, num_fmt_views),
        (grand_aug_posts + grand_sep_posts, num_fmt_int),
    ]

    for col_i, (v, fmt) in enumerate(final_grand_vals, start=1):
        cell = ws.cell(row=row_idx, column=col_i, value=v)
        cell.font = font_grand_total
        cell.number_format = fmt
        cell.alignment = align_center
        cell.fill = fill_grand_total
        cell.border = border_grand_total

    # Set Column Widths for Center Alignment & Readability
    column_widths = {
        1: 18,  # Micro Unit
        2: 15,  # Role / Status
        3: 26,  # Creator Name
        4: 12,  # Channels
        5: 18,  # Aug IG Views
        6: 15,  # Aug IG Reels
        7: 18,  # Sep IG Views
        8: 15,  # Sep IG Reels
        9: 18,  # Aug YT Views
        10: 15, # Aug YT Videos
        11: 18, # Sep YT Views
        12: 15, # Sep YT Videos
        13: 20, # Aug Total Views
        14: 20, # Sep Total Views
        15: 24, # Combined Total Views
        16: 16, # Total Content Items
    }

    for col_idx_num, w in column_widths.items():
        col_letter = get_column_letter(col_idx_num)
        ws.column_dimensions[col_letter].width = w

    # ---------------- SHEET 2: TOP YOUTUBE VIDEOS ----------------
    ws_yt = wb.create_sheet(title="Top YouTube Videos")
    ws_yt.views.sheetView[0].showGridLines = True

    font_link = Font(name="Calibri", size=10, color="1D4ED8", underline="single")

    ws_yt.merge_cells("A1:F2")
    b_yt = ws_yt["A1"]
    b_yt.value = "YouTube Top Videos Report — August & September 2026 (From YouTube App Channel Analytics)"
    b_yt.font = font_title
    b_yt.alignment = align_center
    for r in range(1, 3):
        for c in range(1, 7):
            ws_yt.cell(row=r, column=c).fill = fill_banner

    yt_aug_data = [
        (1, "Sadhguru met Visitors In Nepal Kathmandu", "This_is_last_time_sg", 796866, "2026-08-15", "https://youtu.be/M2D2eRLfhMk"),
        (2, "Sadhguru Comment on YASH as Ravana In Ramayana Mobile", "This_is_last_time_sg", 577034, "2026-08-05", "https://youtu.be/fwldWBv5klg"),
        (3, "⛈️Rainy Day at Isha & Sadhguru Rides 🚴🏻‍♀️ #isha", "This_is_last_time_sg", 545020, "2026-07-31", "https://youtu.be/PliCasQBk30"),
        (4, "Sadhguru Funny Mode 🤣😍 Cute Shahrukh Khan Pose", "This_is_last_time_sg", 491271, "2026-08-16", "https://youtu.be/baA_TWSQZks"),
        (5, "Sadhguru Grown Her Daughter Radhe #sadhguru #radhe", "This_is_last_time_sg", 440394, "2026-08-05", "https://youtu.be/R8ZLmMPvrTc"),
        (6, "ಒಂದು ಅಪರೂಪದ ಅವಕಾಶ.. ನಾಗನಿಗೆ ಬೆಣ್ಣೆ ಸೇವೆ ಅರ್ಪಿಸಿ", "Isha Kannada", 339769, "2026-08-01", "https://youtu.be/0egUGIC3Hzc"),
        (7, "MODERN BABA & BUSINESSMAN Sadhguru's Reality", "This_is_last_time_sg", 295507, "2026-08-26", "https://youtu.be/5N9rtbksVy0"),
        (8, "পড়া মনে রাখার বৈজ্ঞানিক উপায়! Toppers Secret", "জীবন-সমস্যার সমাধান", 218219, "2026-03-09", "https://youtu.be/IUghBabjyEI"),
        (9, "Bageshwar Dham Dhirendera Shastri Met Sadhguru", "This_is_last_time_sg", 208092, "2026-07-21", "https://youtu.be/TnB9Ud-oPaY"),
        (10, "The day Vijji worshipped Sadhguru as God!", "Stillchemy", 207710, "2026-02-26", "https://youtu.be/uMgUZAkS1Gc"),
        (11, "Sadhguru In Flight - Kailash Yatra From Nepal", "This_is_last_time_sg", 190665, "2026-08-15", "https://youtu.be/FSrbqwTa2Wg"),
        (12, "MCQ-তে কখনো ভুল করবে না! ৫টা Scientific Trick", "জীবন-সমস্যার সমাধান", 179518, "2026-04-06", "https://youtu.be/x0LZtPgzNSk"),
        (13, "Ram Charan and Wife with Sadhguru at Home #ramcharan", "This_is_last_time_sg", 174107, "2026-06-06", "https://youtu.be/1lrH0OQaTSo"),
        (14, "🔴LIVE | Sadhguru’s Warning: Do Not Visit Kailash", "Sadhguru Darshan", 168774, "2026-08-30", "https://youtu.be/-mVX4V6aTlU"),
        (15, "Sadhguru returns to Nepal via land route...", "Stillchemy", 145871, "2026-08-30", "https://youtu.be/dBSMfH-cu8c"),
    ]

    yt_sep_data = [
        (1, "Sadhana In Train 🧘🏻‍♀️🔥 Fire Of Sadhana #Sadhguru", "This_is_last_time_sg", 967362, "2025-03-12", "https://youtu.be/VUc5Kvtu9ck"),
        (2, "Sadhguru met Visitors In Nepal Kathmandu", "This_is_last_time_sg", 785284, "2026-08-15", "https://youtu.be/M2D2eRLfhMk"),
        (3, "When the Pursuit of Happiness Comes at the Cost of Peace", "Stillchemy", 453697, "2026-08-31", "https://youtu.be/0G1wrSkMSyI"),
        (4, "Sadhguru says 'he is not so stable…' after the tragedy", "Stillchemy", 437117, "2026-09-13", "https://youtu.be/P7HXt-4Fjnk"),
        (5, "The Nepal Flood Is Warning of Nature | Sadhguru Explains", "Its Meera", 365292, "2026-08-31", "https://youtu.be/JIK-PmMWPBo"),
        (6, "MODERN BABA & BUSINESSMAN Sadhguru's Reality", "This_is_last_time_sg", 344019, "2026-08-26", "https://youtu.be/5N9rtbksVy0"),
        (7, "Sadhguru | The Brutal Truth Behind the Nepal Landslide", "Sadhguru Darshan", 333703, "2026-09-08", "https://youtu.be/43gvG6zWtbA"),
        (8, "Ishanga Praveen Anna's Son l Sadhguru l Nepal", "This_is_last_time_sg", 331314, "2026-08-30", "https://youtu.be/m6JsABwzbNw"),
        (9, "Mastery over 108 Chakras - Tat wale Babaji #Sadhguru", "This_is_last_time_sg", 330259, "2024-10-27", "https://youtu.be/MD_QnTakRdE"),
        (10, "“I Wish We Had Gone There Half an Hour Later…”", "Spiritual Seeker", 291529, "2026-08-31", "https://youtu.be/0MbRa7G8H3A"),
        (11, "Sadhguru's Eye's 🔥 Glasses 🕶️ 🕉️", "This_is_last_time_sg", 287804, "2026-09-01", "https://youtu.be/hSr1x5NVOBA"),
        (12, "Praveen’s son calls for a coordinated rescue system", "Stillchemy", 244739, "2026-08-30", "https://youtu.be/PUcVt49rU1c"),
        (13, "She Thaught \"Sadhguru Is a SCAM\" 😮 #spirituality", "This_is_last_time_sg", 239462, "2026-09-08", "https://youtu.be/29o9D8L-ikQ"),
        (14, "Sadhguru - How Neem Karoli Baba Taught Great Lesson", "Mystics of India", 234082, "2025-01-01", "https://youtu.be/X7AUKlRKsNk"),
        (15, "ಈ 5 ವಿಷಯಗಳನ್ನು ನಿಮ್ಮ ಮುಂದಿನ ಊಟದಲ್ಲಿ ಪ್ರಯತ್ನಿಸಿ", "Isha Kannada", 226946, "2026-09-27", "https://youtu.be/YV8NZxVMNL4"),
    ]

    r_yt = 4
    # August Section
    ws_yt.merge_cells(f"A{r_yt}:F{r_yt}")
    s_aug = ws_yt[f"A{r_yt}"]
    s_aug.value = "🏆 TOP 15 YOUTUBE VIDEOS — AUGUST 2026 (PERIOD VIEWS GAINED)"
    s_aug.font = font_sec_header
    s_aug.fill = fill_yt_header
    s_aug.alignment = align_center
    for c in range(1, 7): ws_yt.cell(row=r_yt, column=c).fill = fill_yt_header; ws_yt.cell(row=r_yt, column=c).border = border_cell
    r_yt += 1

    headers_yt = ["Rank", "Video Title", "YouTube Channel", "Period Views Gained", "Publish Date", "Video URL"]
    for c, h in enumerate(headers_yt, 1):
        cell = ws_yt.cell(row=r_yt, column=c, value=h)
        cell.font = font_tbl_header
        cell.fill = fill_banner
        cell.alignment = align_center
        cell.border = border_cell
    r_yt += 1

    for r_data in yt_aug_data:
        rank, title, ch, v, pdate, url = r_data
        ws_yt.cell(row=r_yt, column=1, value=f"#{rank}").font = font_data_bold
        ws_yt.cell(row=r_yt, column=2, value=title).font = font_data
        ws_yt.cell(row=r_yt, column=3, value=ch).font = font_data_bold
        cell_v = ws_yt.cell(row=r_yt, column=4, value=v)
        cell_v.font = font_winner_views if rank == 1 else font_data
        cell_v.number_format = num_fmt_views
        ws_yt.cell(row=r_yt, column=5, value=pdate).font = font_data
        cell_u = ws_yt.cell(row=r_yt, column=6, value=url)
        cell_u.font = font_link
        cell_u.hyperlink = url

        for c in range(1, 7):
            ws_yt.cell(row=r_yt, column=c).alignment = align_center
            ws_yt.cell(row=r_yt, column=c).border = border_cell
            if rank == 1: ws_yt.cell(row=r_yt, column=c).fill = fill_hl_views
            elif rank % 2 == 1: ws_yt.cell(row=r_yt, column=c).fill = fill_zebra
        r_yt += 1

    r_yt += 2
    # September Section
    ws_yt.merge_cells(f"A{r_yt}:F{r_yt}")
    s_sep = ws_yt[f"A{r_yt}"]
    s_sep.value = "🏆 TOP 15 YOUTUBE VIDEOS — SEPTEMBER 2026 (PERIOD VIEWS GAINED)"
    s_sep.font = font_sec_header
    s_sep.fill = fill_yt_header
    s_sep.alignment = align_center
    for c in range(1, 7): ws_yt.cell(row=r_yt, column=c).fill = fill_yt_header; ws_yt.cell(row=r_yt, column=c).border = border_cell
    r_yt += 1

    for c, h in enumerate(headers_yt, 1):
        cell = ws_yt.cell(row=r_yt, column=c, value=h)
        cell.font = font_tbl_header
        cell.fill = fill_banner
        cell.alignment = align_center
        cell.border = border_cell
    r_yt += 1

    for r_data in yt_sep_data:
        rank, title, ch, v, pdate, url = r_data
        ws_yt.cell(row=r_yt, column=1, value=f"#{rank}").font = font_data_bold
        ws_yt.cell(row=r_yt, column=2, value=title).font = font_data
        ws_yt.cell(row=r_yt, column=3, value=ch).font = font_data_bold
        cell_v = ws_yt.cell(row=r_yt, column=4, value=v)
        cell_v.font = font_winner_views if rank == 1 else font_data
        cell_v.number_format = num_fmt_views
        ws_yt.cell(row=r_yt, column=5, value=pdate).font = font_data
        cell_u = ws_yt.cell(row=r_yt, column=6, value=url)
        cell_u.font = font_link
        cell_u.hyperlink = url

        for c in range(1, 7):
            ws_yt.cell(row=r_yt, column=c).alignment = align_center
            ws_yt.cell(row=r_yt, column=c).border = border_cell
            if rank == 1: ws_yt.cell(row=r_yt, column=c).fill = fill_hl_views
            elif rank % 2 == 1: ws_yt.cell(row=r_yt, column=c).fill = fill_zebra
        r_yt += 1

    yt_widths = {1: 10, 2: 45, 3: 25, 4: 20, 5: 15, 6: 35}
    for col_idx, w in yt_widths.items():
        ws_yt.column_dimensions[get_column_letter(col_idx)].width = w

    # ---------------- SHEET 3: TOP INSTAGRAM REELS ----------------
    ws_ig = wb.create_sheet(title="Top Instagram Reels")
    ws_ig.views.sheetView[0].showGridLines = True

    ws_ig.merge_cells("A1:F2")
    b_ig = ws_ig["A1"]
    b_ig.value = "Instagram Top Reels Report — August & September 2026 (From Scraped Reels Analytics)"
    b_ig.font = font_title
    b_ig.alignment = align_center
    for r in range(1, 3):
        for c in range(1, 7):
            ws_ig.cell(row=r, column=c).fill = fill_banner

    ig_aug_data = [
        (1, "@fantastic_yogi", "Wait, this can actually blow your mind... 🤯👇", 1520776, "2026-08-04", "https://www.instagram.com/p/Dbm5we2NLd0/"),
        (2, "@house_of_pondy", "The true spirit of women’s empowerment was on full display...", 973034, "2026-08-04", "https://www.instagram.com/p/DbnooHwzT9m/"),
        (3, "@timeless_sadhana", "It’s been almost a year on my mind to establish a...", 944724, "2026-08-20", "https://www.instagram.com/p/DcQ0GYexC34/"),
        (4, "@farmer.arun07", "ఈ రోజులలో చాలా మంది జుట్టు రాలిపోతుంది అని బాధపడుతు...", 919855, "2026-08-11", "https://www.instagram.com/p/Db5qgpFCTN8/"),
        (5, "@farmer.arun07", "అందరు ఎలా ఉన్నారు దోస్తులు,నేను బాగున్నా...", 688863, "2026-08-20", "https://www.instagram.com/p/DcPws2Gi-xr/"),
        (6, "@this_is_last_time_sg", "Sadhguru looks too much of a modern Baba & Businessman...", 605785, "2026-08-26", "https://www.instagram.com/p/DcfW39Vpuno/"),
        (7, "@the_taste_of_truth", "She laughed at Pushpaka Vimana and got schooled...", 537914, "2026-08-01", "https://www.instagram.com/p/DbgNWvPzmZV/"),
        (8, "@awakenedbysadhguru", "“A terrible and tragic flash flood has taken place...", 501313, "2026-08-26", "https://www.instagram.com/p/Dcgwcp7swdr/"),
        (9, "@this_is_last_time_sg", "Reality 🔱 Sarabhai Vs Sarabhai Actor Rajesh Kumar...", 489561, "2026-08-24", "https://www.instagram.com/p/DcaFLMXJoP0/"),
        (10, "@krishz_vlogs", "இது வெறும் விளையாட்டு இல்ல… ஒரு ஊరే சேர்ந்து...", 470641, "2026-08-11", "https://www.instagram.com/p/Db5iBRTzdBw/"),
        (11, "@mumbaibappastories", "Trending audio ❌ Dhol Tasha ✅ The most celebrated...", 469842, "2026-08-06", "https://www.instagram.com/p/DbtMhKstvFT/"),
        (12, "@farmer.arun07", "మోకాళ్ళ నొప్పులు నడుము నొప్పులతో బాధపడేవాళ్ళకి...", 346894, "2026-08-18", "https://www.instagram.com/p/DcLgwoMCUtJ/"),
        (13, "@niharshembekar", "आनंद पोटात माज्या माईना रे माईना… 🙏✨ To sing Anan...", 325779, "2026-08-05", "https://www.instagram.com/p/Dbp8feKz71C/"),
        (14, "@farmer.arun07", "పొలాలల్లో పని చేసే రైతులకి ఎంత గానో ఉపయోగపడే...", 314749, "2026-08-09", "https://www.instagram.com/p/Db0dIrjidNP/"),
        (15, "@sadhguru.tamil.quotes", "மண்ணில் பிறந்த மனிதர்கள் யாவரும்...", 310542, "2026-08-14", "https://www.instagram.com/p/Db-e3c3J8Fp/"),
    ]

    ig_sep_data = [
        (1, "@vintiidnani", "No fixed plans, just going with the flow and somehow...", 8501394, "2026-09-16", "https://www.instagram.com/p/DdWExxpBQP1/"),
        (2, "@theindianyogagirl", "During my early age, I didn’t really see a problem...", 3312284, "2026-09-13", "https://www.instagram.com/p/DdOb_cYP4cU/"),
        (3, "@timeless_sadhana", "“A day with Adiyogi”. On the Lap of the Master...", 2407263, "2026-09-14", "https://www.instagram.com/p/DdRNvPxJRjV/"),
        (4, "@blossoming_in_silence", "The longing to experience him #sadhguru #shiva...", 756828, "2026-09-09", "https://www.instagram.com/p/DdEVu7KRguK/"),
        (5, "@blossoming_in_silence", "What is your “higher”? #sadhguru #innerjourney...", 702626, "2026-09-07", "https://www.instagram.com/p/Dc-AqfTN1M9/"),
        (6, "@farmer.arun07", "శంకుపులతో ని కషాయం చేసుకొని తాగితే...", 700677, "2026-09-21", "https://www.instagram.com/p/DdjN5Fwq-8B/"),
        (7, "@blossoming_in_silence", "You stopped looking for strength in others and found...", 609634, "2026-09-16", "https://www.instagram.com/p/DdVQfMjN7pW/"),
        (8, "@theindianyogagirl", "‘Shaucha’ one of the five Niyamas in Yoga...", 592020, "2026-09-22", "https://www.instagram.com/p/Ddl6-MJSYSN/"),
        (9, "@this_is_last_time_sg", "Dhurandhar Actress Sara Arjun shared about Her School...", 516197, "2026-09-11", "https://www.instagram.com/p/DdJuJkVsDJs/"),
        (10, "@yoga_with_divyansh", "Avoiding something that you crave for is one of the...", 486587, "2026-09-22", "https://www.instagram.com/p/DdlKUPcAE0W/"),
        (11, "@sadhguru.archives", "Sadhguru on Conscious Planet & Tree Planting", 460792, "2026-09-16", "https://www.instagram.com/p/DdWfsFZM17v/"),
        (12, "@shisanhi.art", "Sada Shiva... Tale of The Eternal No-Thing.... -x死", 455617, "2026-09-12", "https://www.instagram.com/p/DdLyqLQIwkT/"),
        (13, "@farmer.arun07", "ఉడకపెట్టిన మొక్కజొన్న కంకులు ఎంత మందికి ఇష్టమో...", 445070, "2026-09-18", "https://www.instagram.com/p/DdbdWrECGfn/"),
        (14, "@farzan.dadachanji", "This is the kind of concert I’d choose any day. 🕉️✨", 423929, "2026-09-15", "https://www.instagram.com/p/DdUIq9aK5q5/"),
        (15, "@sadhguru.hindi.gyan", "कर्म का असली रहस्य क्या है? सद्गुरु से जानिए", 412850, "2026-09-05", "https://www.instagram.com/p/Dc3K1jMpc5G/"),
    ]

    r_ig = 4
    # August Section
    ws_ig.merge_cells(f"A{r_ig}:F{r_ig}")
    s_aug_ig = ws_ig[f"A{r_ig}"]
    s_aug_ig.value = "🏆 TOP 15 INSTAGRAM REELS — AUGUST 2026"
    s_aug_ig.font = font_sec_header
    s_aug_ig.fill = fill_ig_header
    s_aug_ig.alignment = align_center
    for c in range(1, 7): ws_ig.cell(row=r_ig, column=c).fill = fill_ig_header; ws_ig.cell(row=r_ig, column=c).border = border_cell
    r_ig += 1

    headers_ig = ["Rank", "Creator Handle", "Caption Snippet", "Video Views", "Publish Date", "Instagram Reel URL"]
    for c, h in enumerate(headers_ig, 1):
        cell = ws_ig.cell(row=r_ig, column=c, value=h)
        cell.font = font_tbl_header
        cell.fill = fill_banner
        cell.alignment = align_center
        cell.border = border_cell
    r_ig += 1

    for r_data in ig_aug_data:
        rank, handle, cap, v, pdate, url = r_data
        ws_ig.cell(row=r_ig, column=1, value=f"#{rank}").font = font_data_bold
        ws_ig.cell(row=r_ig, column=2, value=handle).font = font_data_bold
        ws_ig.cell(row=r_ig, column=3, value=cap).font = font_data
        cell_v = ws_ig.cell(row=r_ig, column=4, value=v)
        cell_v.font = font_winner_views if rank == 1 else font_data
        cell_v.number_format = num_fmt_views
        ws_ig.cell(row=r_ig, column=5, value=pdate).font = font_data
        cell_u = ws_ig.cell(row=r_ig, column=6, value=url)
        cell_u.font = font_link
        cell_u.hyperlink = url

        for c in range(1, 7):
            ws_ig.cell(row=r_ig, column=c).alignment = align_center
            ws_ig.cell(row=r_ig, column=c).border = border_cell
            if rank == 1: ws_ig.cell(row=r_ig, column=c).fill = fill_hl_views
            elif rank % 2 == 1: ws_ig.cell(row=r_ig, column=c).fill = fill_zebra
        r_ig += 1

    r_ig += 2
    # September Section
    ws_ig.merge_cells(f"A{r_ig}:F{r_ig}")
    s_sep_ig = ws_ig[f"A{r_ig}"]
    s_sep_ig.value = "🏆 TOP 15 INSTAGRAM REELS — SEPTEMBER 2026"
    s_sep_ig.font = font_sec_header
    s_sep_ig.fill = fill_ig_header
    s_sep_ig.alignment = align_center
    for c in range(1, 7): ws_ig.cell(row=r_ig, column=c).fill = fill_ig_header; ws_ig.cell(row=r_ig, column=c).border = border_cell
    r_ig += 1

    for c, h in enumerate(headers_ig, 1):
        cell = ws_ig.cell(row=r_ig, column=c, value=h)
        cell.font = font_tbl_header
        cell.fill = fill_banner
        cell.alignment = align_center
        cell.border = border_cell
    r_ig += 1

    for r_data in ig_sep_data:
        rank, handle, cap, v, pdate, url = r_data
        ws_ig.cell(row=r_ig, column=1, value=f"#{rank}").font = font_data_bold
        ws_ig.cell(row=r_ig, column=2, value=handle).font = font_data_bold
        ws_ig.cell(row=r_ig, column=3, value=cap).font = font_data
        cell_v = ws_ig.cell(row=r_ig, column=4, value=v)
        cell_v.font = font_winner_views if rank == 1 else font_data
        cell_v.number_format = num_fmt_views
        ws_ig.cell(row=r_ig, column=5, value=pdate).font = font_data
        cell_u = ws_ig.cell(row=r_ig, column=6, value=url)
        cell_u.font = font_link
        cell_u.hyperlink = url

        for c in range(1, 7):
            ws_ig.cell(row=r_ig, column=c).alignment = align_center
            ws_ig.cell(row=r_ig, column=c).border = border_cell
            if rank == 1: ws_ig.cell(row=r_ig, column=c).fill = fill_hl_views
            elif rank % 2 == 1: ws_ig.cell(row=r_ig, column=c).fill = fill_zebra
        r_ig += 1

    ig_widths = {1: 10, 2: 24, 3: 45, 4: 18, 5: 15, 6: 42}
    for col_idx, w in ig_widths.items():
        ws_ig.column_dimensions[get_column_letter(col_idx)].width = w

    output_path = "/Users/sadhguruwisdomwarrior/.gemini/antigravity/scratch/WisdomWarriors/Wisdom_Warriors_Micro_Units_August_September_2026_Report.xlsx"
    wb.save(output_path)
    print(f"Report saved to {output_path}")

    # Copy to Downloads folder
    dl_path = "/Users/sadhguruwisdomwarrior/Downloads/Wisdom_Warriors_Micro_Units_August_September_2026_Report.xlsx"
    wb.save(dl_path)
    print(f"Report updated in Downloads: {dl_path}")

    # Copy to artifacts path
    artifact_path = "/Users/sadhguruwisdomwarrior/.gemini/antigravity/brain/da87c5c7-cae1-4727-8648-8c431ee707c7/Wisdom_Warriors_Micro_Units_August_September_2026_Report.xlsx"
    wb.save(artifact_path)
    print(f"Copied to artifact path: {artifact_path}")

if __name__ == "__main__":
    asyncio.run(generate_comprehensive_report())
